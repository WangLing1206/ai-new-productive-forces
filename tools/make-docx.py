# -*- coding: utf-8 -*-
"""
生成《智启新质》网站关键亮点说明文档（DOCX）
用法: python tools/make-docx.py
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'deliverables')
os.makedirs(OUT, exist_ok=True)

NAVY = RGBColor(0x0A, 0x2A, 0x5E)
GOLD = RGBColor(0xA8, 0x80, 0x1A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x5A, 0x6B, 0x82)
RED = RGBColor(0xB3, 0x12, 0x1F)

SITE = "https://wangling1206.github.io/ai-new-productive-forces/"

doc = Document()

# ------------------------------------------------------------------ 全局样式
st = doc.styles['Normal']
st.font.name = '微软雅黑'
st.font.size = Pt(10.5)
st.font.color.rgb = INK
st._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(6)

for sec in doc.sections:
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)


def set_font(run, name='微软雅黑', size=10.5, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def heading(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    if level == 1:
        set_font(r, size=15, bold=True, color=NAVY)
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '10')
        bottom.set(qn('w:color'), 'D8AE4E')
        bdr.append(bottom)
        pPr.append(bdr)
    elif level == 2:
        set_font(r, size=12, bold=True, color=GOLD)
    else:
        set_font(r, size=11, bold=True, color=NAVY)
    return p


def para(text, size=10.5, color=INK, bold=False, first_indent=True, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(21)
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold)
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.7)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_font(r1, size=10.5, bold=True, color=NAVY)
    r = p.add_run(text)
    set_font(r, size=10.5)
    return p


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)


def table(headers, rows, widths=None, font_size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_font(r, size=font_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(hdr[i], '0A2A5E')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            r = p.add_run(str(v))
            set_font(r, size=font_size)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text)
    set_font(r, size=9, color=GREY)
    return p


# ================================================================== 封面
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(60)
p.paragraph_format.space_after = Pt(4)
set_font(p.add_run('智 启 新 质'), size=30, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
set_font(p.add_run('人工智能赋能新质生产力 · 全景数据观察'), size=14, color=GOLD)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(30)
set_font(p.add_run('网站关键亮点说明'), size=19, bold=True, color=NAVY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(40)
set_font(p.add_run('AI · NEW PRODUCTIVE FORCES'), size=10, color=GREY)

table(
    ['项目', '内容'],
    [
        ['网站名称', '智启新质 · 人工智能赋能新质生产力全景数据观察'],
        ['在线地址', SITE],
        ['页面数量', '4 个（总览 / 翻页演示 / 数据看板 / 数据来源）'],
        ['可视化结构', '堆叠式滚动叙事 · 翻页式演示 · 响应式数据看板 · 双主题'],
        ['可视化形式', '20 余种（12 类图表 + 8 类自研可视化）'],
        ['数据来源', '24 项权威来源（政府统计 / 研究机构 / 国际组织 / 权威媒体）'],
        ['数据时间跨度', '2019 — 2026 年，数据截至 2026 年 10 月'],
        ['技术栈', 'HTML / CSS / 原生 JavaScript + Apache ECharts 5.5.1（本地化）'],
        ['交付日期', '2026 年 10 月'],
    ],
    widths=[3.4, 12.4], font_size=10
)

doc.add_page_break()

# ================================================================== 一、设计立意
heading('一、设计立意：为什么用"人工智能 × 新质生产力"这个命题')

para('新质生产力是创新起主导作用、具有高科技、高效能、高质量特征、以全要素生产率大幅提升为核心标志的先进生产力质态。'
     '而人工智能是当下唯一同时改变劳动资料（算力与算法）、劳动对象（数据要素）与劳动者（人机协同）的技术变量——'
     '这正是新质生产力"以劳动者、劳动资料、劳动对象及其优化组合的跃升为基本内涵"的当代注脚。')

para('因此，本网站没有停留在"人工智能很厉害"的泛泛表述上，而是建立了一条完整的论证链：'
     '底座（算力与网络）→ 规模（产业与经营主体）→ 效率（实体经济的实测改进）→ 人的维度（就业与人才结构）→ '
     '政策路径（"人工智能+"行动）。每一环都用可核查的数据说话，并逐条标注来源与统计口径。')

heading('网站的三个核心判断', 2)
bullet('人工智能可渗透几乎所有行业，符合新质生产力对"技术革命性突破"的要求。', '它是通用目的技术——')
bullet('算力成为新的劳动资料、数据成为新的劳动对象、人机协同重塑劳动者，三者同时跃升。', '它同时改变三要素——')
bullet('产业规模、效率改进、就业结构三条证据链共同指向全要素生产率的系统性提升。', '它的落点是全要素生产率——')

# ================================================================== 二、结构总览
heading('二、网站结构总览')

table(
    ['页面', '文件', '核心内容', '承担的"结构"要求'],
    [
        ['总览', 'index.html',
         '封面（Canvas 神经网络动态背景）→ 核心论点 → 六章滚动叙事（理论坐标 / 底座跃升 / 规模扩张 / 产业转型 / 人的维度 / 未来路径）→ 数据口径说明 → 结语',
         '堆叠式结构'],
        ['翻页演示', 'deck.html',
         '17 页全屏演示，每页一种可视化；支持键盘、滚轮、触屏、目录跳转、自动播放、全屏',
         '翻页式结构'],
        ['数据看板', 'dashboard.html',
         '8 张核心指标卡 + 12 组图表 + 4 张数据表；支持分类筛选、图表/数据表双视图、打印导出',
         '响应式结构'],
        ['数据来源', 'sources.html',
         '24 项来源清单（含原文链接）、来源可信度分级、5 处口径差异说明、引用格式、技术实现说明',
         '方法论支撑'],
    ],
    widths=[2.0, 2.4, 8.0, 2.6]
)

# ================================================================== 三、四种可视化结构
heading('三、四种可视化结构（对课程要求的逐项落实）')

heading('1. 堆叠式结构 —— 总览页的滚动叙事', 2)
para('这是网站的技术难点与视觉核心。总览页的第二章、第三章采用"滚动叙事（Scrollytelling）"结构：')
bullet('右侧是一块粘性（sticky）仪表盘，内部以"堆叠"方式纵向排列四行指标（算力总规模 / 智能算力 / 5G基站 / 万卡级智算集群）。')
bullet('左侧文字随滚动推进；当阅读到某一指标时，右侧堆叠列表中对应行自动高亮（金色边框 + 位移 + 透明度变化），其余行淡出。')
bullet('同时，堆叠仪表盘下方的详情图会随当前指标切换：算力用渐变柱状图、智能算力用带"口径变更"标注的对比柱、5G基站用面积折线、集群用容量计量条。')
bullet('点击堆叠行可直接跳转到对应步骤，形成双向联动。')
note('技术实现：IntersectionObserver 判定当前步骤 + ECharts setOption 增量切换 + 自研 SVG 迷你走势图，三者在同一容器内协同。')

heading('2. 翻页式结构 —— 17 页全屏演示', 2)
para('翻页演示把整份报告压缩为 17 页，每页一个论点、一种可视化，适合课堂讲解与汇报场景。')
bullet('导航方式：左右方向键 / 上下方向键 / 空格 / PageUp、PageDown / 鼠标滚轮 / 触屏滑动 / 底部按钮 / 目录浮层直接跳转。')
bullet('快捷键：S 呼出目录、P 自动播放（7 秒一页）、F 全屏、Home / End 跳转首尾页。')
bullet('视觉反馈：方向感知的滑动过渡、底部 17 段进度刻度、页码计数、顶部章节标签与页码。')

heading('3. 响应式结构 —— 数据看板', 2)
para('数据看板采用 12 栅格自适应布局，共 12 组图表，按内容权重分配 3 / 4 / 5 / 6 / 7 / 8 栅格宽度，形成专业仪表盘式的非对称版面。')
bullet('分类筛选：全部 / 基础设施 / 产业规模 / 创新与效率 / 就业与人才 / 综合对比，标签上直接显示该分类的图表数量。')
bullet('双视图切换：一键在"图表"与"数据表"之间切换，数据表逐行列出来源机构的原始数值、单位与统计口径。')
bullet('断点适配：1440 / 1024 / 768 / 390px 四档，卡片数与图表高度自动重排；移动端自动折叠为单列并启用汉堡菜单。')
bullet('支持打印 / 导出 PDF（专门的打印样式，隐藏导航与动效）。')

heading('4. 双主题 —— 官方正式的配色体系', 2)
para('全站提供深色与浅色两套主题，配色取政务深蓝 + 赤金 + 中国红点缀，兼顾正式感与科技感；主题选择会被记忆。')
bullet('深色（默认）：深夜蓝底 + 金色高亮 + 青蓝数据色，适合投影与视频演示。')
bullet('浅色：纸白底 + 深蓝标题 + 深金强调，适合打印与阅读。')
bullet('切换主题时，全部 ECharts 图表会同步重绘——通过把 CSS 变量解析为具体色值后再注入 canvas 实现主题联动。')

# ================================================================== 四、可视化形式清单
heading('四、可视化形式清单（20 余种）')

heading('图表类（Apache ECharts，12 类）', 3)
table(
    ['形式', '应用位置', '说明'],
    [
        ['断裂柱状图', '产业规模', '用红色虚线显式标出"统计口径变更"，并保留断点与预测值虚线框'],
        ['对数刻度柱状图', '智能算力 / 私营AI投资', '跨数量级数据用对数轴同图呈现，并注明"对数刻度"'],
        ['双轴柱线复合图', '数字经济规模', '柱表示规模（万亿元）、折线表示占 GDP 比重（%）'],
        ['面积折线图', '5G基站 / 备案大模型', '渐变面积强化增长趋势'],
        ['平滑曲线图', '增长倍数对比', '突出长期趋势'],
        ['堆叠柱状图', '集群与算力设施', '分类叠加对比'],
        ['环形占比图', '生成式AI专利分布', '甜甜圈图呈现中美与其他地区占比'],
        ['仪表盘 Gauge', '生成式AI普及率', '半环仪表呈现 50% 普及率'],
        ['瀑布图', '就业结构变动', '用透明基座柱实现"新增—消失—净增"的瀑布逻辑'],
        ['横向进度条组', '智能制造全球位置', '中国在密度/保有量/安装量三项指标上的相对水平'],
        ['对数横向条形图', '效率提升倍率', '把"提升 X%"统一换算为倍率后横向比较'],
        ['分类散点/条形', '来源类型分布', '统计 24 项来源的机构类型构成'],
    ],
    widths=[3.0, 3.6, 9.2]
)

heading('自研可视化类（原生 SVG / Canvas / DOM，8 类）', 3)
table(
    ['形式', '说明'],
    [
        ['神经网络粒子背景', 'Canvas 绘制，节点自动漂移、按距离连线；鼠标移动时节点与指针产生金色连线，随主题变化'],
        ['迷你面积走势图', '自研 SVG 生成器，用于堆叠仪表盘的每行指标（含对数刻度支持）'],
        ['容量计量条', '自研 SVG，用于无完整时间序列的单值指标'],
        ['政策时间轴', '左右交错布局的双栏时间轴，关键节点（新质生产力提出、政治局集体学习、国发11号文）金色高亮'],
        ['三阶段环形进度', 'SVG stroke-dasharray 描边动画，呈现 2027 / 2030 / 2035 三阶段目标'],
        ['极坐标环形图', '"人工智能+"六大重点行动按极坐标排布，中心为"人工智能+"枢纽'],
        ['效率刻度条', '行业案例卡内的对数刻度进度条'],
        ['数字滚动 / 逐字浮现', 'KPI 数字的缓动计数动画、滚动触发的元素分组揭示'],
    ],
    widths=[3.6, 12.2]
)

# ================================================================== 五、数据可信度
heading('五、数据可信度设计（本网站最核心的差异化亮点）')
para('人工智能领域的数据统计口径极为混乱，同一指标在不同来源间可能出现数倍差距。'
     '本网站选择的是"如实呈现差异"，而不是挑选最好看的数字。这一部分的工作量约占整体的三分之一。')

heading('1. 四条数据处理原则', 2)
bullet('优先使用政府统计公报、国家部委发布口径与官方规划文件原文；所有数据给出机构、报告名称与原文链接。', '只使用可追溯来源：')
bullet('同一指标存在多套口径时分别标注含义，不混用、不拼接；口径变更的序列在图上显式标出断点。', '口径优先于数值：')
bullet('已发生的统计数据与未来测算、规划目标在视觉上严格区分，预测数据一律加"预测"标记。', '区分统计事实与预测：')
bullet('对尚无官方测算的指标不给出单一结论数字，而是并列呈现不同机构的测算区间与分歧。', '不制造精度幻觉：')

heading('2. 五处必须提醒的口径差异', 2)
table(
    ['差异', '具体说明'],
    [
        ['AI 产业规模：新旧口径不可衔接',
         '2019—2023 年为"人工智能核心产业规模"（2023 年 5784 亿元）；2024 年起中国信通院启用新方法测算"人工智能产业规模"（2024 年超 9000 亿元）。强行连线会得出错误的"同比 +55%"结论，本站在图中画出断裂标记。'],
        ['算力规模：三种精度口径并存',
         '工信部 FP32 口径下 2024 年底为 280 EFLOPS，智能算力 90 EFLOPS；2025 年起智能算力改用新精度口径达 1590 EFLOPS。数值相差数倍源于精度标准与统计边界不同。'],
        ['数字经济：全口径与核心产业口径',
         '全口径数字经济规模 2024 年 59.2 万亿元、占 GDP 43.8%（信通院）；数字经济核心产业增加值 14.09 万亿元、占 GDP 10.5%（国家统计局）。两者含义不同。'],
        ['企业数量：AI 企业与"用 AI 的企业"',
         '"人工智能相关企业"（2026 年 6 月超 6600 家）与"开发或应用 AI 的企业"（2024 年同比增长 36%）是两种统计对象。'],
        ['国际对比：投资口径差异显著',
         '斯坦福 HAI 统计的"私营 AI 投资"仅含私募与风险投资，不含政府投入与公开市场融资，与国内投融资总额口径不可直接比较。'],
    ],
    widths=[4.2, 11.6]
)

heading('3. 数据标记体系', 2)
bullet('由官方或研究机构发布的已发生统计数据，含口径说明与年份。', '【统计 / 实测】')
bullet('机构对未来时点的测算或政策文件设定的目标值，不是已发生事实。', '【预测 / 规划目标】')
bullet('该指标存在统计标准变更或多种口径并存，阅读时需注意不可直接比较。', '【口径提示】')
bullet('来自具体企业或地区的实测结果，不代表行业整体水平。', '【案例样本】')

# ================================================================== 六、内容亮点
heading('六、内容与数据亮点')

table(
    ['维度', '关键数据', '来源'],
    [
        ['产业规模', '2025 年人工智能产业规模超 1.2 万亿元，同比 +40%；2027 年预计达 2 万亿元', '中国信通院 / 工信部'],
        ['算力底座', '算力总规模四年从 140 增至 280 EFLOPS；2026 年 6 月智能算力 2185 EFLOPS', '工业和信息化部'],
        ['连接底座', '5G 基站从 13 万个增至 483.8 万个（2025 年底）', '工信部《通信业统计公报》'],
        ['模型供给', '累计备案生成式 AI 模型 748 款，2025 年新增备案 446 款', '国家互联网信息办公室'],
        ['用户扩散', '生成式 AI 用户规模突破 7 亿人，普及率超 50%', 'CNNIC'],
        ['经营主体', '人工智能相关企业超 6600 家（2026 年 6 月），全球占比约 15%', '中国信通院'],
        ['产业母体', '数字经济规模 59.2 万亿元，占 GDP 43.8%', '中国信通院'],
        ['智能制造', '工业机器人安装量占全球 54%，保有量 202.7 万台占 43.5%，密度居全球第三', 'IFR'],
        ['创新产出', '2024 年全球新公开生成式 AI 专利中，中国 2.7 万件、占 61.5%', 'WIPO 检索方法'],
        ['效率改进', '电网仿真效率提升 30 倍以上；药物研发成本降低近一半；单吨钢材成本降 97 元', '人民网等公开报道'],
        ['就业结构', '到 2030 年全球岗位净增约 7800 万个（新增 1.7 亿、消失 9200 万）', '世界经济论坛（预测）'],
        ['人才供需', '中国 AI 人才缺口超 500 万人，供求比约 1:10', '人社部 / 智联招聘'],
        ['政策目标', '2027 年智能终端与智能体应用普及率超 70%，2030 年超 90%，2035 年步入智能社会', '国发〔2025〕11 号'],
    ],
    widths=[2.3, 9.3, 4.2]
)

# ================================================================== 七、技术亮点
heading('七、技术与工程亮点')

bullet('全站为 HTML / CSS / 原生 JavaScript，无框架、无构建步骤、无后端依赖。', '纯静态、零构建：')
bullet('图表库 Apache ECharts 5.5.1 已本地化部署，不加载任何外部 CDN、第三方字体或分析脚本；可离线打开、可长期归档。', '零外部依赖：')
bullet('三个页面共用同一份 data.js，任何数据修改一处即全站同步，杜绝多页面数字不一致。', '单一数据源：')
bullet('ECharts 通过 canvas 渲染，无法识别 CSS 变量；本站通过"CSS 变量 → 具体色值"的运行时解析层，实现了图表与主题的完整联动。', '主题联动：')
bullet('图片级优化：迷你走势图为自研 SVG（无请求、矢量清晰）；看板图表采用 IntersectionObserver 懒加载，避免首屏一次性渲染 12 张图表。', '性能设计：')
bullet('无障碍与可用性：键盘完整导航、语义化标签、aria 标注、尊重系统"减少动态效果"偏好、打印样式适配。', '可访问性：')
bullet('提供 5 个 Playwright 自动化脚本（整页截图、逐页截图、响应式断点、浅色主题、线上巡检），用于持续验收。', '工程化验收：')

# ================================================================== 八、演示视频
heading('八、演示视频说明')
table(
    ['项目', '内容'],
    [
        ['文件名', '智启新质-演示视频.mp4'],
        ['时长 / 规格', '4 分 31 秒 · 1920×1080 · 25fps · H.264（yuv420p）· 约 52 MB'],
        ['内容结构', '15 个场景，覆盖封面、核心论点、六章滚动叙事、数据口径说明、数据来源页、翻页演示、数据看板交互与双主题、结语'],
        ['字幕', '已烧录中文字幕（微软雅黑 44px，白字深描边），另附独立字幕文件 subtitles.srt 可自行替换'],
        ['音轨', '原片为静音立体声音轨，直接添加音轨配音即可；讲解词与逐句时间码见《讲解词与时间码》'],
        ['制作方式', 'Playwright 脚本化巡览录制（含虚拟鼠标指针指示交互）→ 黑场精准定位起点 → FFmpeg 裁切、烧录字幕、片尾淡出、H.264 编码'],
    ],
    widths=[3.2, 12.6]
)

# ================================================================== 九、交付清单
heading('九、交付清单')
table(
    ['序号', '交付物', '说明'],
    [
        ['1', '网页链接', SITE],
        ['2', '网站源代码（ZIP）', '仅包含网站源码：4 个 HTML 页面、assets（CSS / JS / ECharts 本地副本）、tools（验收脚本）、README、LICENSE'],
        ['3', '演示视频', 'MP4，4 分 31 秒，含讲解词、烧录字幕与时间码，可自行配音'],
        ['4', '本文档', '网站关键亮点说明'],
    ],
    widths=[1.2, 4.2, 10.4]
)

# ================================================================== 附录
doc.add_page_break()
heading('附录：数据来源清单（24 项）')

SOURCES = [
    ('中国政府网', '《关于深入实施"人工智能+"行动的意见》(国发〔2025〕11号)', '2025-08-26'),
    ('中国政府网', '《政府工作报告》全文（2024 年、2025 年）', '2024 / 2025'),
    ('人民网 / 新华社', '习近平在中共中央政治局第十一次集体学习时的讲话', '2024-02-01'),
    ('工业和信息化部', '《通信业统计公报》(2019—2024)', '逐年'),
    ('工业和信息化部', '人工智能产业规模、智能算力、企业数量发布数据', '2025—2026'),
    ('国家互联网信息办公室', '生成式人工智能服务备案信息公告', '2024—2026'),
    ('国家统计局', '数字经济核心产业增加值核算数据', '2024'),
    ('国家知识产权局', '发明专利有效量及专利密集型产业统计', '2025'),
    ('中国互联网络信息中心 CNNIC', '生成式人工智能用户规模调查报告', '2026'),
    ('中国信息通信研究院', '人工智能产业规模测算 / 算力规模报告', '2024—2026'),
    ('中国信息通信研究院', '《中国数字经济发展研究报告(2025年)》', '2026-03'),
    ('中国电子学会', '《人工智能发展白皮书》', '2021'),
    ('世界经济论坛 WEF', '《未来就业报告2025》', '2025'),
    ('世界经济论坛 WEF', '全球灯塔工厂网络报告', '2025—2026'),
    ('国际机器人联合会 IFR', '《世界机器人》报告', '2025'),
    ('斯坦福大学 HAI', '《AI Index 2025 / 2026》', '2025 / 2026'),
    ('普华永道 PwC', '《Sizing the Prize》', '2017'),
    ('麦肯锡 McKinsey', '《The economic potential of generative AI》', '2023'),
    ('国际货币基金组织 IMF', 'AI Will Transform the Global Economy', '2024-01'),
    ('经济合作与发展组织 OECD', 'Artificial Intelligence 专题研究', '2025'),
    ('世界知识产权组织 WIPO', '生成式人工智能专利检索分析', '2025'),
    ('人民网 / 人民日报', '各行业人工智能应用成效系列报道', '2024—2026'),
    ('智联招聘', '《2025年人工智能产业人才发展报告》', '2025-10'),
    ('清华大学互联网产业研究院', 'AI 与全要素生产率学术观点汇编', '2025'),
]
table(['序号', '机构', '文件 / 报告名称', '时间'],
      [[str(i + 1), a, b, c] for i, (a, b, c) in enumerate(SOURCES)],
      widths=[1.2, 4.0, 8.4, 2.2])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
r = p.add_run('说明：本网站为数据可视化研究作品，所有原始数据的版权归原发布机构所有；'
              '数据的整理、可视化呈现与口径注释为原创工作。报告中涉及的国际机构测算为其研究结论，'
              '涉及未来时点的数据均为预测，不构成任何投资建议。')
set_font(r, size=9, color=GREY)

out_file = os.path.join(OUT, '智启新质-网站关键亮点说明.docx')
doc.save(out_file)
print('已生成：' + out_file)
